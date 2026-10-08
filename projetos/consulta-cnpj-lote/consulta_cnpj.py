# -*- coding: utf-8 -*-
"""Consulta CNPJ em lote e enriquece uma planilha Excel."""

import json
import re
import sys
import time
from datetime import datetime
from pathlib import Path

try:
    import requests
    from openpyxl import load_workbook
except ImportError as exc:
    print("Biblioteca ausente:", exc)
    print("Execute EXECUTAR.bat para instalar as dependências.")
    input("Pressione ENTER para sair...")
    sys.exit(1)

API = "https://brasilapi.com.br/api/cnpj/v1/{}"
USER_AGENT = "ConsultaCNPJEmLote/1.0"
PAUSA_ENTRE_CONSULTAS = 2.5
MAX_TENTATIVAS = 4
TIMEOUT = 25

CAMPOS = [
    ("Logradouro", "logradouro"),
    ("Número", "numero"),
    ("Complemento", "complemento"),
    ("Bairro", "bairro"),
    ("Cidade", "municipio"),
    ("UF", "uf"),
    ("CEP", "cep"),
]


def limpar_cnpj(valor):
    if valor is None:
        return ""
    return re.sub(r"[^0-9A-Z]", "", str(valor).strip().upper())


def encontrar_excel():
    arquivos = [
        p for p in Path(".").glob("*.xlsx")
        if not p.name.startswith("~$")
        and p.name != "relatorio_CNPJ_enriquecido.xlsx"
    ]
    if not arquivos:
        raise FileNotFoundError(
            "Nenhum arquivo .xlsx de entrada foi encontrado nesta pasta."
        )
    return max(arquivos, key=lambda p: p.stat().st_size)


def encontrar_aba(wb):
    if "Resultado da consulta" in wb.sheetnames:
        return wb["Resultado da consulta"]
    return wb.active


def encontrar_coluna_cnpj(ws):
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 10)):
        for cell in row:
            if cell.value is not None and "cnpj" in str(cell.value).strip().lower():
                return cell.column
    raise ValueError("Não encontrei uma coluna contendo 'CNPJ'.")


def consultar(cnpj, session):
    for tentativa in range(1, MAX_TENTATIVAS + 1):
        try:
            resposta = session.get(
                API.format(cnpj),
                headers={"Accept": "application/json"},
                timeout=TIMEOUT,
            )

            if resposta.status_code == 200:
                return "OK", resposta.json()

            if resposta.status_code == 404:
                return "NÃO ENCONTRADO", {}

            if resposta.status_code in (429, 500, 502, 503, 504):
                espera = 5 * tentativa
                print(
                    f"    HTTP {resposta.status_code}; "
                    f"aguardando {espera}s "
                    f"(tentativa {tentativa}/{MAX_TENTATIVAS})"
                )
                time.sleep(espera)
                continue

            return f"ERRO HTTP {resposta.status_code}", {}

        except requests.RequestException as exc:
            espera = 5 * tentativa
            print(
                f"    Falha de conexão; aguardando {espera}s "
                f"(tentativa {tentativa}/{MAX_TENTATIVAS})"
            )
            time.sleep(espera)
            ultimo_erro = str(exc)

    return f"FALHA: {ultimo_erro[:100]}", {}


def main():
    print("=" * 70)
    print(" CONSULTA DE CNPJ EM LOTE")
    print("=" * 70)

    entrada = encontrar_excel()
    print(f"Arquivo de entrada: {entrada.name}")

    wb = load_workbook(entrada)
    ws = encontrar_aba(wb)
    coluna_cnpj = encontrar_coluna_cnpj(ws)

    cabecalho = {
        str(c.value).strip(): c.column
        for c in ws[1]
        if c.value is not None
    }

    colunas = {}
    for nome, _ in CAMPOS:
        if nome in cabecalho:
            colunas[nome] = cabecalho[nome]
        else:
            colunas[nome] = ws.max_column + 1
            ws.cell(1, colunas[nome], nome)

    for nome in ("Status CNPJ", "Data da consulta"):
        if nome in cabecalho:
            colunas[nome] = cabecalho[nome]
        else:
            colunas[nome] = ws.max_column + 1
            ws.cell(1, colunas[nome], nome)

    cache_path = Path("cache_cnpj.json")
    if cache_path.exists():
        try:
            cache = json.loads(cache_path.read_text(encoding="utf-8"))
        except Exception:
            cache = {}
    else:
        cache = {}

    cnpjs = []
    for row in range(2, ws.max_row + 1):
        cnpj = limpar_cnpj(ws.cell(row, coluna_cnpj).value)
        if cnpj and cnpj not in cnpjs:
            cnpjs.append(cnpj)

    print(f"Linhas do relatório: {ws.max_row - 1}")
    print(f"CNPJs únicos: {len(cnpjs)}")
    print(f"CNPJs já no cache: {sum(c in cache for c in cnpjs)}")

    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT})

    for indice, cnpj in enumerate(cnpjs, 1):
        if cnpj in cache:
            print(f"[{indice}/{len(cnpjs)}] {cnpj} — cache")
            continue

        print(f"[{indice}/{len(cnpjs)}] Consultando {cnpj}...")
        status, dados = consultar(cnpj, session)

        cache[cnpj] = {
            "status": status,
            "data": {chave: dados.get(chave) for _, chave in CAMPOS}
            if dados else {},
            "consultado_em": datetime.now().isoformat(timespec="seconds"),
        }

        cache_path.write_text(
            json.dumps(cache, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

        time.sleep(PAUSA_ENTRE_CONSULTAS)

    for row in range(2, ws.max_row + 1):
        cnpj = limpar_cnpj(ws.cell(row, coluna_cnpj).value)
        item = cache.get(cnpj, {"status": "SEM CNPJ", "data": {}})
        dados = item.get("data", {})

        for nome, chave in CAMPOS:
            valor = dados.get(chave, "")
            if chave == "cep" and valor:
                valor = re.sub(r"\D", "", str(valor))
            ws.cell(row, colunas[nome], valor or "")

        ws.cell(row, colunas["Status CNPJ"], item.get("status", ""))
        ws.cell(
            row,
            colunas["Data da consulta"],
            item.get("consultado_em", ""),
        )

    saida = Path("relatorio_CNPJ_enriquecido.xlsx")
    wb.save(saida)

    ok = sum(cache.get(c, {}).get("status") == "OK" for c in cnpjs)
    nao = sum(
        cache.get(c, {}).get("status") == "NÃO ENCONTRADO"
        for c in cnpjs
    )

    print()
    print("=" * 70)
    print(" CONCLUÍDO")
    print("=" * 70)
    print(f"Consultas OK: {ok}")
    print(f"Não encontrados: {nao}")
    print(f"Outros erros/falhas: {len(cnpjs) - ok - nao}")
    print(f"Arquivo gerado: {saida.resolve()}")
    print("A planilha original não foi alterada.")
    input("Pressione ENTER para fechar...")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print()
        print("ERRO:", exc)
        print("A planilha original não foi alterada.")
        input("Pressione ENTER para fechar...")
